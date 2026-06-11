#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
§3.3: ссылки на рисунки админки из гл. 2 (без дублирования скринов).

py fix_vkr5_ch3_admin_fig_refs.py
"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document

from thesis_paragraph_format import style_body_paragraph_content

ROOT = Path(__file__).resolve().parent
TARGET = ROOT / "ВКР5.docx"
BACKUP_SUFFIX = ".bak_before_admin_refs"

CH3_3_TITLE = "Техническая документация администратора"

INTRO = (
    "Ниже приведена краткая техническая документация для сопровождения production-стенда. "
    "Внешний вид экранов административной зоны описан в разделе 2 (рисунки 12, 30–34). "
    "Авторизация — через Supabase Auth (email и пароль); роль берётся из поля `profiles.role` "
    "и определяет состав меню `AdminSidebar`. Вход в админ-зону: `/admin` (далее редирект по роли)."
)

# (заголовок роли, абзацы текста)
ROLE_TEXT: list[tuple[str, list[str]]] = [
    (
        "Супер-администратор (super_admin)",
        [
            "Полный доступ ко всем разделам `/admin`: дашборд со статистикой (рисунок 12), CRUD туров и медиа (S3) "
            "(рисунок 31), бронирования, комнаты туров, отзывы, жалобы в чатах и на гидов (рисунки 33–34), "
            "чат поддержки, пользователи (смена роли, бан/разбан) (рисунок 32), апелляции на бан, выдача достижений.",
            "Может блокировать любого пользователя, кроме ограничений логики UI для других супер-админов. "
            "Изменения ролей и банов синхронизируются клиенту через Pusher (`admin-sync`).",
        ],
    ),
    (
        "Администратор туров (tour_admin)",
        [
            "Доступ к `/admin` с дашбордом «Панель управления» (без редиректа на moderator-dashboard). "
            "В меню: туры (рисунок 31), бронирования, комнаты туров, жалобы в чатах туров и на гидов "
            "(рисунок 34), отзывы (рисунок 33), апелляции на бан, выдача достижений; пункты «Чат поддержки» "
            "и «Пользователи» скрыты.",
            "Не может менять глобальные роли (`/api/admin/users/role` — только super_admin) и не открывает "
            "реестр пользователей (рисунок 32): прямой переход на `/admin/users` перенаправляет на `/admin`. "
            "При обращении к API чата поддержки `/api/admin/support/*` сервер возвращает 403.",
        ],
    ),
    (
        "Модератор поддержки (support_admin)",
        [
            "После входа — редирект на `/admin/moderator-dashboard`. Доступны: отзывы (публикация/снятие) "
            "(рисунок 33), жалобы в чатах туров и на гидов (рисунок 34), комнаты туров (просмотр), чат поддержки, "
            "пользователи (бан с ограничениями: нельзя заблокировать super_admin, tour_admin и другого support_admin) "
            "(рисунок 32), апелляции.",
            "Не создаёт и не редактирует туры (рисунок 31), не управляет бронированиями как tour_admin.",
        ],
    ),
    (
        "Гид (guide)",
        [
            "Вход на `/admin/guide-dashboard`. В меню: «Мои туры», «Выдача достижений», комнаты закреплённых туров. "
            "Может выдать участнику экспертное достижение в комнате тура или в `/admin/award-achievements` "
            "(рисунок 30).",
            "Не видит чужие туры, пользователей, бронирования и модераторские очереди (рисунки 31–34).",
        ],
    ),
]

CLOSING = (
    "При смене пароля или роли в Supabase обновите таблицу 7 и повторите smoke-тест: "
    "вход, состав меню и отказ в доступе к чужим разделам (`403` или редирект)."
)


def backup(path: Path) -> Path:
    bak = path.with_suffix(path.suffix + BACKUP_SUFFIX)
    shutil.copy2(path, bak)
    return bak


def find_section_33(doc: Document) -> tuple[int, int]:
    start = end = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t == CH3_3_TITLE:
            start = i
            continue
        if start is not None and t == "Заключение":
            end = i
            break
    if start is None:
        raise RuntimeError(f"«{CH3_3_TITLE}» не найден")
    return start, end or len(doc.paragraphs)


def set_para(p, text: str) -> None:
    style_body_paragraph_content(p, text)


def process(path: Path) -> dict[str, int]:
    doc = Document(str(path))
    start, end = find_section_33(doc)
    stats = {"intro": 0, "roles": 0, "closing": 0}

    paras = doc.paragraphs[start:end]
    # intro — первый абзац после заголовка
    for p in paras[1:]:
        t = p.text.strip()
        if t.startswith("Ниже приведена"):
            set_para(p, INTRO)
            stats["intro"] = 1
            break

    role_idx = 0
    i = 1
    while i < len(paras):
        t = paras[i].text.strip()
        if t == "Заключение" or t.startswith("При смене пароля"):
            set_para(paras[i], CLOSING)
            stats["closing"] = 1
            break
        if role_idx < len(ROLE_TEXT) and t == ROLE_TEXT[role_idx][0]:
            _, blocks = ROLE_TEXT[role_idx]
            j = 0
            k = i + 1
            while k < len(paras) and j < len(blocks):
                nt = paras[k].text.strip()
                if not nt or nt == ROLE_TEXT[role_idx][0]:
                    k += 1
                    continue
                if role_idx + 1 < len(ROLE_TEXT) and nt == ROLE_TEXT[role_idx + 1][0]:
                    break
                if nt.startswith("Таблица ") or nt.startswith("При смене"):
                    break
                set_para(paras[k], blocks[j])
                j += 1
                k += 1
            stats["roles"] += 1
            role_idx += 1
            i = k
            continue
        i += 1

    doc.save(str(path))
    return stats


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("doc", nargs="?", default=str(TARGET))
    parser.add_argument("--no-backup", action="store_true")
    args = parser.parse_args()
    path = Path(args.doc).resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    if not args.no_backup:
        print("Backup:", backup(path).name)
    stats = process(path)
    print("fix_vkr5_ch3_admin_fig_refs:", path.name, stats)


if __name__ == "__main__":
    main()
