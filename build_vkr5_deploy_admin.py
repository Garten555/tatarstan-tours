#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
ВКР4 → бэкап + ВКР5:
  — §3.2: два рисунка деплоя (SSH-команды на VPS, smoke-тест);
  — §3.3: таблица ролей и текст, без скринов админки.

py build_vkr5_deploy_admin.py
"""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from fix_vkr5_ch3_subsection_lists import (
    CH3_3_SUB_NUM_ID,
    ensure_subsection_numbering,
    fix_ch3_head_and_sections,
    fix_ch3_subsections,
    find_ch3_bounds,
)
from integrate_admin_panels_docx import insert_after
from thesis_document import THESIS_DOC
from thesis_paragraph_format import (
    apply_thesis_body_format,
    set_run_font,
    style_body_paragraph_content,
    style_figure_caption_content,
    style_figure_placeholder_content,
)

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "ВКР4.docx"
BACKUP = ROOT / "ВКР4.docx.bak_before_v5"
TARGET = ROOT / "ВКР5.docx"

OUTLINE_NUM_ID = 2
FONT_BODY = "Times New Roman"
FIG_START = 42

CH3_3_TITLE = "Техническая документация администратора"

# Тестовые учётные записи (заменить на реальные перед сдачей)
TEST_ACCOUNTS: list[tuple[str, str, str, str]] = [
    ("super_admin", "admin.super@vkr-tours.test", "SuperAdmin2026!", "/admin"),
    ("tour_admin", "tours.admin@vkr-tours.test", "TourAdmin2026!", "/admin"),
    ("support_admin", "support.mod@vkr-tours.test", "SupportMod2026!", "/admin/moderator-dashboard"),
    ("guide", "guide.kazan@vkr-tours.test", "GuideKazan2026!", "/admin/guide-dashboard"),
]

DEPLOY_SCREENSHOTS: list[tuple[str, str, str]] = [
    (
        "Повторные выкладки выполнялись вручную по SSH",
        "[МЕСТО ДЛЯ РИСУНКА: SSH-сессия на VPS — git pull, npm ci, npm run build, pm2 restart tatarstan-tours]",
        "Рисунок {n} — Обновление production-версии на VPS (команды в SSH-сессии)",
    ),
    (
        "После каждой выкладки выполнялся smoke-тест",
        "[МЕСТО ДЛЯ РИСУНКА: главная страница или каталог на production-URL после деплоя]",
        "Рисунок {n} — Проверка работоспособности сайта после выкладки",
    ),
]

ROLE_DOCS: list[tuple[str, list[str]]] = [
    (
        "Супер-администратор (super_admin)",
        [
            "Полный доступ ко всем разделам `/admin`: дашборд со статистикой (рисунок 12), CRUD туров и медиа (S3) "
            "(рисунок 31), бронирования, комнаты туров, отзывы, жалобы в чатах и на гидов (рисунки 33–34), "
            "чат поддержки, пользователи (смена роли, бан/разбан) (рисунок 32), апелляции на бан, выдача достижений.",
            "Может блокировать любого пользователя, кроме ограничений логики UI для других супер-админов. "
            "Изменения ролей и банов синхронизируются клиенту через Pusher (`admin-sync`).",
        ],
    ),
    (
        "Администратор туров (tour_admin)",
        [
            "Доступ к `/admin` с дашбордом «Панель управления» (без редиректа на moderator-dashboard). "
            "В меню: туры (рисунок 31), бронирования, комнаты туров, жалобы в чатах туров и на гидов "
            "(рисунок 34), отзывы (рисунок 33), апелляции на бан, выдача достижений; пункты «Чат поддержки» "
            "и «Пользователи» скрыты.",
            "Не может менять глобальные роли (`/api/admin/users/role` — только super_admin) и не открывает "
            "реестр пользователей (рисунок 32): прямой переход на `/admin/users` перенаправляет на `/admin`. "
            "При обращении к API чата поддержки `/api/admin/support/*` сервер возвращает 403.",
        ],
    ),
    (
        "Модератор поддержки (support_admin)",
        [
            "После входа — редирект на `/admin/moderator-dashboard`. Доступны: отзывы (публикация/снятие) "
            "(рисунок 33), жалобы в чатах туров и на гидов (рисунок 34), комнаты туров (просмотр), чат поддержки, "
            "пользователи (бан с ограничениями: нельзя заблокировать super_admin, tour_admin и другого support_admin) "
            "(рисунок 32), апелляции.",
            "Не создаёт и не редактирует туры (рисунок 31), не управляет бронированиями как tour_admin.",
        ],
    ),
    (
        "Гид (guide)",
        [
            "Вход на `/admin/guide-dashboard`. В меню: «Мои туры», «Выдача достижений», комнаты закреплённых туров. "
            "Может выдать участнику экспертное достижение в комнате тура или в `/admin/award-achievements` "
            "(рисунок 30).",
            "Не видит чужие туры, пользователей, бронирования и модераторские очереди (рисунки 31–34).",
        ],
    ),
]


def backup_and_copy() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Нет исходника: {SOURCE}")
    shutil.copy2(SOURCE, BACKUP)
    shutil.copy2(SOURCE, TARGET)


def set_numbering(p: Paragraph, ilvl: int, num_id: int = OUTLINE_NUM_ID) -> None:
    pPr = p._element.get_or_add_pPr()
    numPr = pPr.find(qn("w:numPr"))
    if numPr is None:
        numPr = OxmlElement("w:numPr")
        pPr.insert(0, numPr)
    else:
        for child in list(numPr):
            numPr.remove(child)
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    numPr.append(ilvl_el)
    numPr.append(num_id_el)


def style_heading(p: Paragraph, text: str, ilvl: int, num_id: int = OUTLINE_NUM_ID) -> None:
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    set_run_font(run, name=FONT_BODY, size=14, bold=True, italic=False)
    try:
        p.style = "List Paragraph"
    except KeyError:
        pass
    set_numbering(p, ilvl, num_id)


def insert_figure_after(anchor: Paragraph, placeholder: str, caption: str) -> Paragraph:
    ph = insert_after(anchor, placeholder)
    style_figure_placeholder_content(ph, placeholder)
    cap = insert_after(ph, caption)
    style_figure_caption_content(cap, caption)
    return cap


def find_paragraph(doc: Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise RuntimeError(f"Абзац не найден: {prefix[:50]!r}…")


def find_conclusion(doc: Document) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip() == "Заключение":
            return p
    raise RuntimeError("«Заключение» не найдено")


def add_deploy_screenshots(doc: Document) -> int:
    n = FIG_START
    count = 0
    for prefix, ph, cap_tpl in DEPLOY_SCREENSHOTS:
        anchor = find_paragraph(doc, prefix)
        cap = cap_tpl.format(n=n)
        insert_figure_after(anchor, ph, cap)
        n += 1
        count += 1
    return count


def insert_section_33(doc: Document) -> dict[str, int]:
    conclusion = find_conclusion(doc)

    heading_el = OxmlElement("w:p")
    conclusion._element.addprevious(heading_el)
    tail = Paragraph(heading_el, conclusion._parent)
    style_heading(tail, CH3_3_TITLE, 1)

    intro = (
        "Ниже приведена краткая техническая документация для сопровождения production-стенда. "
        "Внешний вид экранов административной зоны описан в разделе 2 (рисунки 12, 30–34). "
        "Авторизация — через Supabase Auth (email и пароль); роль берётся из поля `profiles.role` "
        "и определяет состав меню `AdminSidebar`. Вход в админ-зону: `/admin` (далее редирект по роли)."
    )
    tail = insert_after(tail, intro)
    style_body_paragraph_content(tail, intro)

    note = (
        "Таблица 7 — Тестовые учётные записи административных ролей "
        "(заменить на фактические перед эксплуатацией)"
    )
    tail = insert_after(tail, note)
    apply_thesis_body_format(tail)
    for r in tail.runs:
        set_run_font(r, name=FONT_BODY, size=14, bold=True)

    table = doc.add_table(rows=1 + len(TEST_ACCOUNTS), cols=4)
    table.style = "Table Grid"
    headers = ("Роль", "E-mail (логин)", "Пароль (тест)", "URL входа")
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = h
        for par in cell.paragraphs:
            for run in par.runs:
                run.bold = True
    for i, (role, email, pwd, url) in enumerate(TEST_ACCOUNTS, start=1):
        cells = table.rows[i].cells
        cells[0].text = role
        cells[1].text = email
        cells[2].text = pwd
        cells[3].text = url
    tbl_el = table._tbl
    doc.element.body.remove(tbl_el)
    tail._element.addnext(tbl_el)
    gap_el = OxmlElement("w:p")
    tbl_el.addnext(gap_el)
    tail = Paragraph(gap_el, tail._parent)

    roles_added = 0
    for title, paras in ROLE_DOCS:
        hel = OxmlElement("w:p")
        tail._element.addnext(hel)
        rh = Paragraph(hel, tail._parent)
        style_heading(rh, title, 1, CH3_3_SUB_NUM_ID)
        tail = rh
        for text in paras:
            tail = insert_after(tail, text)
            style_body_paragraph_content(tail, text)
        roles_added += 1

    closing = (
        "При смене пароля или роли в Supabase обновите таблицу 7 и повторите smoke-тест: "
        "вход, состав меню и отказ в доступе к чужим разделам (`403` или редирект)."
    )
    tail = insert_after(tail, closing)
    style_body_paragraph_content(tail, closing)

    return {"roles": roles_added, "figures_from": FIG_START}



def process() -> dict[str, int]:
    backup_and_copy()
    ensure_subsection_numbering(TARGET)
    doc = Document(str(TARGET))
    stats = {
        "deploy_figures": add_deploy_screenshots(doc),
        **insert_section_33(doc),
    }
    doc.save(str(TARGET))
    doc = Document(str(TARGET))
    start, end = find_ch3_bounds(doc)
    stats["ch3_lists"] = {
        "headings": fix_ch3_head_and_sections(doc, start, end),
        **fix_ch3_subsections(doc, start, end),
    }
    doc.save(str(TARGET))
    return stats


def main() -> None:
    global SOURCE, TARGET, BACKUP
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", default=str(SOURCE))
    parser.add_argument("--target", default=str(TARGET))
    args = parser.parse_args()
    SOURCE = Path(args.source).resolve()
    TARGET = Path(args.target).resolve()
    BACKUP = SOURCE.with_suffix(SOURCE.suffix + ".bak_before_v5")
    stats = process()
    print("Backup:", BACKUP.name)
    print("Created:", TARGET.name, stats)


if __name__ == "__main__":
    main()
